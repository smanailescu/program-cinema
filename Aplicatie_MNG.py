import streamlit as st
import mysql.connector
import random
import time
import io
from datetime import datetime, timedelta, date
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

st.set_page_config(layout="wide", page_title="Panou Control Manager")

if "manager_logat" not in st.session_state:
    st.session_state.manager_logat = False
if "id_angajat_de_editat" not in st.session_state:
    st.session_state.id_angajat_de_editat = None

TRADUCERE_ZILE = {
    'Monday': 'Luni', 'Tuesday': 'Marti', 'Wednesday': 'Miercuri',
    'Thursday': 'Joi', 'Friday': 'Vineri', 'Saturday': 'Sambata', 'Sunday': 'Duminica'
}


def get_conexiune():
    # Dacă aplicația este urcată online, ia datele din Cloud:
    try:
        if "DB_HOST" in st.secrets:
            return mysql.connector.connect(
                host=st.secrets["DB_HOST"],
                port=int(st.secrets["DB_PORT"]),
                user=st.secrets["DB_USER"],
                password=st.secrets["DB_PASSWORD"],
                database=st.secrets["DB_NAME"]
            )
    except Exception:
        pass

    # Altfel, dacă o rulezi de pe PC-ul tău, se conectează automat la XAMPP:
    return mysql.connector.connect(
        host="localhost",
        user="root",
        password="",
        database="nume_angajati"
    )


def normalizeaza_text(txt):
    if not txt:
        return ""
    return (
        txt.strip().lower()
        .replace("ă", "a").replace("â", "a")
        .replace("î", "i").replace("ș", "s").replace("ş", "s")
        .replace("ț", "t").replace("ţ", "t")
    )


# Verifică STRICT dacă abilitățile angajatului corespund cu departamentul cerut în șablon
def are_abilitate(dep_necesar, deps_angajat_str):
    if not deps_angajat_str:
        return False
    deps = [normalizeaza_text(d) for d in deps_angajat_str.split(",") if d.strip()]
    dep_n = normalizeaza_text(dep_necesar)

    if "vip" in dep_n:
        return any("vip" in d for d in deps)
    elif "bar" in dep_n:
        return any("bar" in d for d in deps)
    elif "plasator" in dep_n:
        return any("plasator" in d and "vip" not in d for d in deps)
    elif "cafe" in dep_n:
        return any("cafe" in d for d in deps)
    elif "bilete" in dep_n or "casier" in dep_n or "box" in dep_n:
        return any(k in d for d in deps for k in ["bilete", "casier", "box"])
    return dep_n in deps
def este_tura_dimineata(ora_start_td):
    start_h = ora_start_td.total_seconds() / 3600.0
    return start_h <= 12.0
def este_tura_seara(ora_start_td, ora_end_td):
    start_h = ora_start_td.total_seconds() / 3600.0
    end_h = ora_end_td.total_seconds() / 3600.0
    if end_h <= start_h:
        end_h += 24.0
    return start_h >= 16.0 or end_h >= 22.0

def potrivire_tura_specifica(pref_tura, ora_start_td, ora_end_td):
    if not pref_tura:
        return False
    pref = normalizeaza_text(pref_tura)

    start_h = ora_start_td.total_seconds() / 3600.0
    end_h = ora_end_td.total_seconds() / 3600.0
    if end_h <= start_h:
        end_h += 24.0

    if pref == "dimineata":
        return start_h <= 12.0
    elif pref == "middle":
        return 12.0 < start_h <= 17.0 and end_h <= 23.5
    elif pref in ["seara", "inchidere", "close"]:
        return start_h >= 16.0 or end_h >= 24.0
    return False

def genereaza_excel_program(date_zile, zile_pe_rand):
    wb = Workbook()
    ws = wb.active
    ws.title = "Program Saptamanal"

    fill_titlu_zi = PatternFill(start_color="262730", end_color="262730", fill_type="solid")
    fill_header = PatternFill(start_color="F4B084", end_color="F4B084", fill_type="solid")
    fill_verde = PatternFill(start_color="00B050", end_color="00B050", fill_type="solid")
    fill_galben = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    fill_rosu = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
    fill_alb = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    font_titlu_zi = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    font_bold_negru = Font(name="Calibri", size=10, bold=True, color="000000")
    align_center = Alignment(horizontal="center", vertical="center")

    thin_side = Side(border_style="thin", color="000000")
    med_side = Side(border_style="medium", color="000000")
    border_celula = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    border_header = Border(left=thin_side, right=thin_side, top=med_side, bottom=med_side)

    rand_start_grup = 2

    for idx_start in range(0, len(date_zile), zile_pe_rand):
        grup = date_zile[idx_start: idx_start + zile_pe_rand]
        max_rand_in_grup = rand_start_grup

        for col_idx, info_zi in enumerate(grup):
            col_baza = col_idx * 5 + 1
            rand_curent = rand_start_grup

            ws.column_dimensions[get_column_letter(col_baza)].width = 18
            ws.column_dimensions[get_column_letter(col_baza + 1)].width = 9
            ws.column_dimensions[get_column_letter(col_baza + 2)].width = 9
            ws.column_dimensions[get_column_letter(col_baza + 3)].width = 24
            ws.column_dimensions[get_column_letter(col_baza + 4)].width = 4

            ws.merge_cells(
                start_row=rand_curent, start_column=col_baza,
                end_row=rand_curent, end_column=col_baza + 3
            )
            titlu_text = f"{info_zi['zi_ro'].upper()} — {info_zi['data_str']}"
            for c in range(col_baza, col_baza + 4):
                cel = ws.cell(row=rand_curent, column=c)
                if c == col_baza:
                    cel.value = titlu_text
                cel.fill = fill_titlu_zi
                cel.font = font_titlu_zi
                cel.alignment = align_center
                cel.border = border_header
            ws.row_dimensions[rand_curent].height = 24
            rand_curent += 1

            headers = ["Departament", "Start", "End", "Angajat"]
            for offset, h_text in enumerate(headers):
                cel = ws.cell(row=rand_curent, column=col_baza + offset, value=h_text)
                cel.fill = fill_header
                cel.font = font_bold_negru
                cel.alignment = align_center
                cel.border = border_header
            ws.row_dimensions[rand_curent].height = 20
            rand_curent += 1

            for rand_tura in info_zi["randuri"]:
                dep, start_str, end_str, nume_ang, mod = rand_tura

                if mod == "Verde":
                    fill_tura = fill_verde
                elif mod == "Galben":
                    fill_tura = fill_galben
                else:
                    fill_tura = fill_rosu

                valori = [dep, start_str, end_str, nume_ang]
                for offset, val in enumerate(valori):
                    cel = ws.cell(row=rand_curent, column=col_baza + offset, value=val)
                    cel.font = font_bold_negru
                    cel.alignment = align_center
                    cel.border = border_celula
                    cel.fill = fill_alb if offset == 3 else fill_tura

                ws.row_dimensions[rand_curent].height = 18
                rand_curent += 1

            if rand_curent > max_rand_in_grup:
                max_rand_in_grup = rand_curent

        rand_start_grup = max_rand_in_grup + 2

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


if not st.session_state.manager_logat:
    st.title("👔 Acces Manager")
    parola_introdusa = st.text_input("Parolă:", type="password")
    if st.button("Intră în cont"):
        if parola_introdusa == "cinema2026":
            st.session_state.manager_logat = True
            st.rerun()
        else:
            st.error("❌ Parolă incorectă!")

else:
    with st.sidebar:
        st.write("👔 **Mod Manager Activ**")
        st.divider()
        if st.button("🚪 Ieși din cont"):
            st.session_state.manager_logat = False
            st.rerun()

    try:
        con = get_conexiune()
        cursor = con.cursor()

        tab1, tab2, tab3, tab4 = st.tabs(["📅 Programul Final", "🤖 Generare Automată", "👥 Echipa", "➕ Adaugă Angajat"])

        with tab1:
            st.title("📅 Programul Întregii Săptămâni")

            cursor.execute("SELECT DISTINCT data_zi FROM program_final ORDER BY data_zi ASC")
            zile_generate = cursor.fetchall()

            if zile_generate:
                date_zile_procesate = []
                for zi_tuple in zile_generate:
                    zi_selectata = zi_tuple[0]
                    nume_zi_sapt = zi_selectata.strftime('%A')
                    zi_ro = TRADUCERE_ZILE.get(nume_zi_sapt, 'Luni')

                    cursor.execute("""
                        SELECT zona, departament, ora_start, ora_end, mod_trafic, necesar_oameni
                        FROM sabloane_necesar
                        WHERE zi_saptamana = %s
                        ORDER BY FIELD(departament, 'VIP Host', 'VIP Plasator', 'Cafe', 'Bilete', 'Bar Mare',
                                       'Bar IMAX', 'Plasator 1-9', 'Plasator 10-15', 'Plasator 16-20',
                                       'Plasator IMAX'),
                                 FIELD(mod_trafic, 'Verde', 'Galben', 'Rosu'),
                                 ora_start ASC
                    """, (zi_ro,))
                    sabloane_zi = cursor.fetchall()

                    cursor.execute("""
                        SELECT pf.departament, pf.ora_start, pf.ora_end, pf.mod_trafic, a.nume_angajat
                        FROM program_final pf
                        LEFT JOIN angajati a ON pf.id_angajat = a.id_angajat
                        WHERE pf.data_zi = %s
                    """, (zi_selectata,))
                    alocati_zi = cursor.fetchall()

                    alocati_map = {}
                    for al in alocati_zi:
                        cheie_al = (al[0], al[1], al[2], al[3])
                        if cheie_al not in alocati_map:
                            alocati_map[cheie_al] = []
                        alocati_map[cheie_al].append(al[4] if al[4] else "-")

                    randuri_zi = []
                    for sablon in sabloane_zi:
                        zona, dep, o_start, o_end, mod, necesar = (
                            sablon[0], sablon[1], sablon[2], sablon[3], sablon[4], sablon[5]
                        )
                        cheie_s = (dep, o_start, o_end, mod)
                        start_str = (datetime.min + o_start).time().strftime('%H:%M')
                        end_str = (datetime.min + o_end).time().strftime('%H:%M')

                        if cheie_s in alocati_map:
                            lista_nume = alocati_map[cheie_s]
                            for _ in range(necesar):
                                nume_angajat = lista_nume.pop(0) if lista_nume else "-"
                                randuri_zi.append((dep, start_str, end_str, nume_angajat, mod))
                        else:
                            for _ in range(necesar):
                                randuri_zi.append((dep, start_str, end_str, "", mod))

                    date_zile_procesate.append({
                        "zi_selectata": zi_selectata,
                        "zi_ro": zi_ro,
                        "data_str": zi_selectata.strftime('%d.%m.%Y'),
                        "randuri": randuri_zi
                    })

                col_act1, col_act2, col_act3 = st.columns([1.1, 1.1, 1.8])

                with col_act3:
                    zile_pe_rand = st.radio(
                        "📐 Câte zile vrei să vezi grupate pe un rând?",
                        options=[2, 3, 4],
                        index=1,
                        horizontal=True
                    )

                with col_act1:
                    if st.button("🗑️ Resetează / Șterge Programul", type="primary", use_container_width=True):
                        cursor.execute("TRUNCATE TABLE program_final")
                        con.commit()
                        st.success("Programul a fost șters!")
                        time.sleep(1)
                        st.rerun()

                with col_act2:
                    excel_bytes = genereaza_excel_program(date_zile_procesate, zile_pe_rand)
                    prima_zi_str = date_zile_procesate[0]["zi_selectata"].strftime('%d_%m_%Y')
                    st.download_button(
                        label="📥 Descarcă Excel (.xlsx)",
                        data=excel_bytes,
                        file_name=f"Program_Cinema_{prima_zi_str}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        use_container_width=True
                    )

                st.write("---")

                for idx_start in range(0, len(date_zile_procesate), zile_pe_rand):
                    grup_zile = date_zile_procesate[idx_start: idx_start + zile_pe_rand]
                    coloane_grid = st.columns(zile_pe_rand, gap="medium")

                    for col_idx, info_zi in enumerate(grup_zile):
                        with coloane_grid[col_idx]:
                            if info_zi["randuri"]:
                                html_table = "<div style='margin-bottom: 25px; box-shadow: 0 2px 6px rgba(0,0,0,0.15);'>"
                                html_table += "<table style='width:100%; border-collapse: collapse; text-align: center; font-family: sans-serif; font-size: 12.5px; color: black; border: 2px solid black;'>"
                                html_table += "<thead>"
                                html_table += "<tr style='background-color: #262730; color: white; font-size: 14px; font-weight: bold; border: 2px solid black;'>"
                                html_table += f"<th colspan='4' style='padding: 8px; text-transform: uppercase; letter-spacing: 0.5px;'>📅 {info_zi['zi_ro']} — {info_zi['data_str']}</th>"
                                html_table += "</tr>"
                                html_table += "<tr style='background-color: #f4b084; color: black; font-weight: bold; border: 2px solid black;'>"
                                html_table += "<th style='border: 1px solid black; padding: 5px;'>Departament</th>"
                                html_table += "<th style='border: 1px solid black; padding: 5px;'>Start</th>"
                                html_table += "<th style='border: 1px solid black; padding: 5px;'>End</th>"
                                html_table += "<th style='border: 1px solid black; padding: 5px; width: 38%;'>Angajat</th>"
                                html_table += "</tr>"
                                html_table += "</thead>"
                                html_table += "<tbody>"

                                for rand_tura in info_zi["randuri"]:
                                    dep, start_str, end_str, nume_angajat, mod = rand_tura

                                    if mod == 'Verde':
                                        bg_color = "#00b050"
                                    elif mod == 'Galben':
                                        bg_color = "#ffff00"
                                    else:
                                        bg_color = "#ff0000"

                                    html_table += f"<tr style='background-color: {bg_color}; color: black; font-weight: bold; border: 1px solid black;'>"
                                    html_table += f"<td style='border: 1px solid black; padding: 4px; white-space: nowrap;'>{dep}</td>"
                                    html_table += f"<td style='border: 1px solid black; padding: 4px;'>{start_str}</td>"
                                    html_table += f"<td style='border: 1px solid black; padding: 4px;'>{end_str}</td>"
                                    html_table += f"<td style='border: 1px solid black; padding: 4px; background-color: white; color: black;'>{nume_angajat}</td>"
                                    html_table += "</tr>"

                                html_table += "</tbody></table></div>"
                                st.markdown(html_table, unsafe_allow_html=True)
                            else:
                                st.warning(f"Nu există șabloane pentru {info_zi['zi_ro']}.")
            else:
                st.info("Niciun program generat încă. Mergi la tab-ul 'Generare Automată' pentru a crea programul săptămânii.")


        with tab2:
            st.title("🤖 Generare Program Inteligent")
            col1, col2, col3 = st.columns(3)
            moduri = ["Verde (Bază)", "Galben (Mediu)", "Roșu (Aglomerat)"]

            mod_vip = col1.selectbox("Mod VIP:", moduri)
            mod_imax = col2.selectbox("Mod IMAX:", moduri)
            mod_cinema = col3.selectbox("Mod Cinema:", moduri)

    
            azi = date.today()
            zile_pana_la_joi = (3 - azi.weekday()) % 7
            if zile_pana_la_joi == 0:
                zile_pana_la_joi = 7
            joi_implicit = azi + timedelta(days=zile_pana_la_joi)

            data_start = st.date_input("Data de început a săptămânii (Joi):", value=joi_implicit)
            sfarsit_sapt = data_start + timedelta(days=6)

            if data_start.weekday() != 3:
                st.warning(
                    f"⚠️ Atenție: Data selectată ({data_start.strftime('%d.%m.%Y')}) este "
                    f"**{TRADUCERE_ZILE[data_start.strftime('%A')]}**, nu Joi!"
                )

            st.write("---")
            st.subheader("🗂️ Gestionare Preferințe Angajați")

            cursor.execute("SELECT MIN(data_zi), MAX(data_zi), COUNT(*) FROM cand_pot_lucra")
            interval_pref = cursor.fetchone()
            if interval_pref and interval_pref[0]:
                st.info(
                    f"📌 Preferințe totale în baza de date: **{interval_pref[2]} opțiuni** "
                    f"(perioada **{interval_pref[0].strftime('%d.%m.%Y')} — {interval_pref[1].strftime('%d.%m.%Y')}**)"
                )
            else:
                st.warning("📭 Momentan nu există nicio preferință salvată în baza de date.")

            col_pref1, col_pref2 = st.columns(2)
            with col_pref1:
                if st.button(
                    f"🧹 Șterge preferințele vechi (înainte de {data_start.strftime('%d.%m.%Y')})",
                    use_container_width=True
                ):
                    cursor.execute("DELETE FROM cand_pot_lucra WHERE data_zi < %s", (data_start,))
                    sters = cursor.rowcount
                    con.commit()
                    st.success(f"✅ Au fost șterse {sters} preferințe din săptămânile trecute!")
                    time.sleep(1.2)
                    st.rerun()

            with col_pref2:
                if st.button("🗑️ Șterge TOATE preferințele (Resetare completă)", use_container_width=True):
                    cursor.execute("TRUNCATE TABLE cand_pot_lucra")
                    con.commit()
                    st.success("✅ Toate preferințele au fost șterse! Angajații pot introduce acum opțiunile actualizate.")
                    time.sleep(1.2)
                    st.rerun()

            with st.expander(
                f"👀 Vezi preferințele actualizate pentru săptămâna selectată "
                f"({data_start.strftime('%d.%m.%Y')} — {sfarsit_sapt.strftime('%d.%m.%Y')})"
            ):
                cursor.execute("""
                    SELECT a.nume_angajat, c.data_zi, c.tura
                    FROM cand_pot_lucra c
                    JOIN angajati a ON c.id_angajat = a.id_angajat
                    WHERE c.data_zi BETWEEN %s AND %s
                    ORDER BY c.data_zi ASC, a.nume_angajat ASC
                """, (data_start, sfarsit_sapt))
                pref_sapt = cursor.fetchall()

                if pref_sapt:
                    tabel_pref = [
                        {
                            "Angajat": r[0],
                            "Ziua": f"{TRADUCERE_ZILE.get(r[1].strftime('%A'), '')} ({r[1].strftime('%d.%m.%Y')})",
                            "Preferință Tură": r[2]
                        }
                        for r in pref_sapt
                    ]
                    st.dataframe(tabel_pref, use_container_width=True)
                else:
                    st.info("Nicio preferință introdusă încă pentru această săptămână.")

            st.write("---")

            if st.button("🚀 Generează Programul", use_container_width=True, type="primary"):
                mapare_mod = {
                    "Verde (Bază)": ["Verde"],
                    "Galben (Mediu)": ["Verde", "Galben"],
                    "Roșu (Aglomerat)": ["Verde", "Galben", "Rosu"]
                }

                cursor.execute("TRUNCATE TABLE program_final")

                cursor.execute("""
                    SELECT a.id_angajat, a.nume_angajat, a.gen, a.rating, GROUP_CONCAT(DISTINCT ad.departament)
                    FROM angajati a
                    JOIN angajat_departament ad ON a.id_angajat = ad.id_angajat
                    GROUP BY a.id_angajat
                """)
                toti_angajatii = cursor.fetchall()
                angajati_dict = {ang[0]: ang for ang in toti_angajatii}

                nr_abilitati = {
                    ang[0]: len([d for d in (ang[4] or "").split(",") if d.strip()])
                    for ang in toti_angajatii
                }

                cursor.execute("SELECT id_angajat, data_zi, tura FROM cand_pot_lucra")
                toate_preferintele = cursor.fetchall()
                pref_dict = {}
                for p in toate_preferintele:
                    cheie = (p[0], p[1])
                    if cheie not in pref_dict:
                        pref_dict[cheie] = []
                    if p[2]:
                        pref_dict[cheie].append(normalizeaza_text(p[2]))

                istoric_saptamanal = {ang[0]: 0 for ang in toti_angajatii}
                lucrat_seara_zi = {}

                for i in range(7):
                    data_curenta = data_start + timedelta(days=i)
                    data_ieri = data_curenta - timedelta(days=1)
                    nume_zi = TRADUCERE_ZILE.get(data_curenta.strftime('%A'), 'Luni')

                    lucrat_seara_zi[data_curenta] = set()
                    angajati_seara_ieri = lucrat_seara_zi.get(data_ieri, set())

                    filtre_vip = "','".join(mapare_mod[mod_vip])
                    filtre_imax = "','".join(mapare_mod[mod_imax])
                    filtre_cinema = "','".join(mapare_mod[mod_cinema])


                    cursor.execute(f"""
                        SELECT id_necesar, departament, ora_start, ora_end, necesar_oameni, zona, mod_trafic 
                        FROM sabloane_necesar 
                        WHERE zi_saptamana = '{nume_zi}' AND (
                            (zona = 'VIP' AND mod_trafic IN ('{filtre_vip}')) OR
                            (zona = 'IMAX' AND mod_trafic IN ('{filtre_imax}')) OR
                            (zona = 'Cinema' AND mod_trafic IN ('{filtre_cinema}'))
                        )
                        ORDER BY FIELD(departament, 'VIP Host', 'VIP Plasator', 'Cafe', 'Bilete',
                                       'Bar IMAX', 'Bar Mare', 'Plasator IMAX', 'Plasator 1-9',
                                       'Plasator 10-15', 'Plasator 16-20'),
                                 ora_start ASC
                    """)
                    ture_necesare = cursor.fetchall()

                    sloturi_zi = []
                    for tura in ture_necesare:
                        dep_necesar = tura[1]
                        ora_s, ora_e = tura[2], tura[3]
                        oameni_necesari = tura[4]
                        mod_tura = tura[6]
                        for _ in range(oameni_necesari):
                            sloturi_zi.append({
                                "dep": dep_necesar,
                                "ora_s": ora_s,
                                "ora_e": ora_e,
                                "mod": mod_tura,
                                "id_alocat": None
                            })

        
                    angajati_folositi_azi = set()

                    def aloca_angajat(slot, id_ales):
                        slot["id_alocat"] = id_ales
                        angajati_folositi_azi.add(id_ales)
                        istoric_saptamanal[id_ales] += 1
                        if este_tura_seara(slot["ora_s"], slot["ora_e"]):
                            lucrat_seara_zi[data_curenta].add(id_ales)

     
                    for slot in sloturi_zi:
                        if slot["id_alocat"] is not None:
                            continue
                        este_dim = este_tura_dimineata(slot["ora_s"])
                        candidati = []
                        for ang in toti_angajatii:
                            id_ang, deps = ang[0], ang[4]
                            if id_ang in angajati_folositi_azi or istoric_saptamanal[id_ang] >= 5:
                                continue
                            if este_dim and id_ang in angajati_seara_ieri:
                                continue
                            if not are_abilitate(slot["dep"], deps):
                                continue
                            prefs = pref_dict.get((id_ang, data_curenta), [])
                            if any(potrivire_tura_specifica(p, slot["ora_s"], slot["ora_e"]) for p in prefs):
                                candidati.append(ang)

                        if candidati:
                            candidati.sort(key=lambda a: (
                                istoric_saptamanal[a[0]] >= 2,
                                istoric_saptamanal[a[0]],
                                -float(a[3] or 0),
                                nr_abilitati[a[0]]
                            ))
                            aloca_angajat(slot, candidati[0][0])

           
                    for slot in sloturi_zi:
                        if slot["id_alocat"] is not None:
                            continue
                        este_dim = este_tura_dimineata(slot["ora_s"])
                        candidati = []
                        for ang in toti_angajatii:
                            id_ang, deps = ang[0], ang[4]
                            if id_ang in angajati_folositi_azi or istoric_saptamanal[id_ang] >= 5:
                                continue
                            if este_dim and id_ang in angajati_seara_ieri:
                                continue
                            if not are_abilitate(slot["dep"], deps):
                                continue
                            prefs = pref_dict.get((id_ang, data_curenta), [])
                            if any(p in ["oricand", "toata ziua", "all"] for p in prefs):
                                candidati.append(ang)

                        if candidati:
                            candidati.sort(key=lambda a: (
                                istoric_saptamanal[a[0]] >= 2,
                                istoric_saptamanal[a[0]],
                                -float(a[3] or 0),
                                nr_abilitati[a[0]]
                            ))
                            aloca_angajat(slot, candidati[0][0])

        
                    for slot in sloturi_zi:
                        if slot["id_alocat"] is not None:
                            continue
                        este_dim = este_tura_dimineata(slot["ora_s"])
                        candidati = []
                        for ang in toti_angajatii:
                            id_ang, deps = ang[0], ang[4]
                            if id_ang in angajati_folositi_azi or istoric_saptamanal[id_ang] >= 5:
                                continue
                            if este_dim and id_ang in angajati_seara_ieri:
                                continue
                            if not are_abilitate(slot["dep"], deps):
                                continue
                            prefs = pref_dict.get((id_ang, data_curenta), [])
                            if not prefs:
                                status_disp = 1
                            elif all(p in ["liber", "indisponibil", "nu pot"] for p in prefs):
                                status_disp = 2
                            else:
                                status_disp = 0
                            candidati.append((ang, status_disp))

                        if candidati:
                            candidati.sort(key=lambda x: (
                                x[1],
                                istoric_saptamanal[x[0][0]] >= 2,
                                float(x[0][3] or 0),
                                istoric_saptamanal[x[0][0]],
                                nr_abilitati[x[0][0]]
                            ))
                            aloca_angajat(slot, candidati[0][0][0])

           
                    for slot in sloturi_zi:
                        if slot["id_alocat"] is not None:
                            continue
                        este_dim = este_tura_dimineata(slot["ora_s"])
                        candidati = []
                        for ang in toti_angajatii:
                            id_ang, deps = ang[0], ang[4]
                            if id_ang in angajati_folositi_azi:
                                continue  
                            if este_dim and id_ang in angajati_seara_ieri:
                                continue
                            if not are_abilitate(slot["dep"], deps):
                                continue
                            candidati.append(ang)

                        if candidati:
                            candidati.sort(key=lambda a: (
                                float(a[3] or 0),
                                istoric_saptamanal[a[0]],
                                nr_abilitati[a[0]]
                            ))
                            aloca_angajat(slot, candidati[0][0])

         
                    for slot_gol in sloturi_zi:
                        if slot_gol["id_alocat"] is not None:
                            continue
                        este_dim_gol = este_tura_dimineata(slot_gol["ora_s"])

                        for slot_ocupat in sloturi_zi:
                            id_ocupat = slot_ocupat["id_alocat"]
                            if id_ocupat is None:
                                continue
                            ang_ocupat = angajati_dict[id_ocupat]

                          
                            if not are_abilitate(slot_gol["dep"], ang_ocupat[4]):
                                continue
                            if este_dim_gol and id_ocupat in angajati_seara_ieri:
                                continue

                            
                            este_dim_ocupat = este_tura_dimineata(slot_ocupat["ora_s"])
                            inlocuitori = []
                            for ang_liber in toti_angajatii:
                                id_liber, deps_liber = ang_liber[0], ang_liber[4]
                                if id_liber in angajati_folositi_azi:
                                    continue  
                                if este_dim_ocupat and id_liber in angajati_seara_ieri:
                                    continue
                                if not are_abilitate(slot_ocupat["dep"], deps_liber):
                                    continue
                                inlocuitori.append(ang_liber)

                            if inlocuitori:
                                inlocuitori.sort(key=lambda a: (
                                    istoric_saptamanal[a[0]] >= 5,
                                    istoric_saptamanal[a[0]] >= 2,
                                    float(a[3] or 0),
                                    istoric_saptamanal[a[0]]
                                ))
                                id_inlocuitor = inlocuitori[0][0]

                                
                                if este_tura_seara(slot_ocupat["ora_s"], slot_ocupat["ora_e"]):
                                    lucrat_seara_zi[data_curenta].discard(id_ocupat)

                                slot_gol["id_alocat"] = id_ocupat
                                if este_tura_seara(slot_gol["ora_s"], slot_gol["ora_e"]):
                                    lucrat_seara_zi[data_curenta].add(id_ocupat)

                                slot_ocupat["id_alocat"] = id_inlocuitor
                                angajati_folositi_azi.add(id_inlocuitor)
                                istoric_saptamanal[id_inlocuitor] += 1
                                if este_tura_seara(slot_ocupat["ora_s"], slot_ocupat["ora_e"]):
                                    lucrat_seara_zi[data_curenta].add(id_inlocuitor)
                                break

                   
                    for slot in sloturi_zi:
                        cursor.execute(
                            "INSERT INTO program_final (id_angajat, data_zi, departament, ora_start, ora_end, mod_trafic) VALUES (%s, %s, %s, %s, %s, %s)",
                            (slot["id_alocat"], data_curenta, slot["dep"], slot["ora_s"], slot["ora_e"], slot["mod"])
                        )

                con.commit()
                st.success("✅ Programul a fost generat! Fiecare angajat are cel mult o tură pe zi și doar pe departamentele cunoscute.")
                time.sleep(1.5)
                st.rerun()

   
        departamente_existente = ["Bar", "Plasator", "VIP", "Cafe", "Bilete"]

        with tab3:
            st.title("👥 Gestionare Echipa")

            cursor.execute("SELECT COUNT(id_angajat) FROM angajati")
            total_angajati = cursor.fetchone()[0]
            st.write(f"Număr total de angajați activi în sistem: **{total_angajati}**")

            with st.expander("📊 Vezi statistici ture angajați (Săptămâna curentă)"):
                cursor.execute("""
                    SELECT a.nume_angajat, pf.departament
                    FROM program_final pf
                    JOIN angajati a ON pf.id_angajat = a.id_angajat
                """)
                ture_alocate = cursor.fetchall()

                if ture_alocate:
                    statistici = {}
                    for nume, dep in ture_alocate:
                        if nume not in statistici:
                            statistici[nume] = {}
                        if dep not in statistici[nume]:
                            statistici[nume][dep] = 0
                        statistici[nume][dep] += 1

                    tabel_statistici = []
                    for nume, deps in statistici.items():
                        total_ture = sum(deps.values())
                        detalii_ture = ", ".join([f"{d} ({c}x)" for d, c in deps.items()])
                        tabel_statistici.append({"Angajat": nume, "Total Ture": total_ture, "Repartizare": detalii_ture})

                    tabel_statistici = sorted(tabel_statistici, key=lambda x: x["Total Ture"], reverse=True)
                    st.dataframe(tabel_statistici, use_container_width=True)
                else:
                    st.info("Nu există ture alocate momentan. Generează programul din tab-ul 'Generare Automată'.")

            st.write("---")

            col_stanga, col_dreapta = st.columns([1, 1.2], gap="large")
            with col_stanga:
                dep_ales = st.selectbox("Filtrează după abilitate:", departamente_existente)
                if dep_ales == "Bilete":
                    cursor.execute("""
                        SELECT DISTINCT a.id_angajat, a.nume_angajat, a.adresa_email, a.rating, a.gen
                        FROM angajati a
                        JOIN angajat_departament ad ON a.id_angajat = ad.id_angajat
                        WHERE ad.departament IN ('Bilete', 'Casier', 'Box')
                    """)
                else:
                    cursor.execute("""
                        SELECT DISTINCT a.id_angajat, a.nume_angajat, a.adresa_email, a.rating, a.gen
                        FROM angajati a
                        JOIN angajat_departament ad ON a.id_angajat = ad.id_angajat
                        WHERE ad.departament LIKE %s
                    """, (f"%{dep_ales}%",))
                angajati_filtrati = cursor.fetchall()

                if angajati_filtrati:
                    st.write("---")
                    for rand in angajati_filtrati:
                        id_ang, nume, email, rating, gen = rand[0], rand[1], rand[2], rand[3], rand[4] if rand[4] else "Nespecificat"
                        c_detalii, c_buton = st.columns([4, 1])
                        c_detalii.markdown(f"**{nume}** ({gen})  \n⭐ {rating} | ✉️ {email}")
                        if c_buton.button("✏️", key=f"edit_{id_ang}"):
                            st.session_state.id_angajat_de_editat = id_ang
                            st.rerun()
                        st.write("---")
                else:
                    st.info("Niciun angajat alocat pentru această abilitate.")

            with col_dreapta:
                if st.session_state.id_angajat_de_editat is not None:
                    id_ales = st.session_state.id_angajat_de_editat
                    cursor.execute("SELECT nume_angajat, adresa_email, rating, gen FROM angajati WHERE id_angajat = %s", (id_ales,))
                    date_ang = cursor.fetchone()
                    if date_ang:
                        cursor.execute("SELECT departament FROM angajat_departament WHERE id_angajat = %s", (id_ales,))
                        deps_cur_raw = [r[0] for r in cursor.fetchall()]
                        deps_cur = ["Bilete" if d in ["Casier", "Box"] else d for d in deps_cur_raw]

                        st.subheader(f"Editează: {date_ang[0]}")
                        nou_email = st.text_input("Email:", value=date_ang[1])
                        nou_gen = st.selectbox("Gen:", ["M", "F"], index=0 if date_ang[3] == "M" else 1)
                        nou_rating = st.number_input("Rating (Max 5.0):", min_value=1.0, max_value=5.0, value=float(date_ang[2]), step=0.1)
                        nou_deps = st.multiselect("Abilități (Departamente):", departamente_existente, default=list(set([d for d in deps_cur if d in departamente_existente])))

                        c1, c2 = st.columns(2)
                        if c1.button("💾 Salvează"):
                            cursor.execute("UPDATE angajati SET adresa_email = %s, rating = %s, gen = %s WHERE id_angajat = %s", (nou_email, nou_rating, nou_gen, id_ales))
                            cursor.execute("DELETE FROM angajat_departament WHERE id_angajat = %s", (id_ales,))
                            for dep in nou_deps:
                                cursor.execute("INSERT INTO angajat_departament (id_angajat, departament) VALUES (%s, %s)", (id_ales, dep))
                            con.commit()
                            st.session_state.id_angajat_de_editat = None
                            st.rerun()
                        if c2.button("❌ Șterge"):
                            cursor.execute("DELETE FROM cand_pot_lucra WHERE id_angajat = %s", (id_ales,))
                            cursor.execute("DELETE FROM angajat_departament WHERE id_angajat = %s", (id_ales,))
                            cursor.execute("DELETE FROM program_final WHERE id_angajat = %s", (id_ales,))
                            cursor.execute("DELETE FROM angajati WHERE id_angajat = %s", (id_ales,))
                            con.commit()
                            st.session_state.id_angajat_de_editat = None
                            st.rerun()

    
        with tab4:
            st.title("➕ Adaugă Angajat Nou")
            n_nume = st.text_input("Nume complet:")
            n_email = st.text_input("Adresa de email:")
            n_gen = st.selectbox("Gen:", ["M", "F"], key="gen_nou")
            n_rating = st.number_input("Rating inițial (1 la 5):", min_value=1.0, max_value=5.0, value=5.0, step=0.1)
            deps_alese = st.multiselect("Abilități (Departamente):", departamente_existente, key="deps_nou")

            if st.button("💾 Înregistrează Angajatul"):
                if n_nume and n_email and deps_alese:
                    nou_id = random.randint(1000, 99999)
                    cursor.execute("INSERT INTO angajati (id_angajat, nume_angajat, adresa_email, rating, gen) VALUES (%s, %s, %s, %s, %s)", (nou_id, n_nume, n_email, n_rating, n_gen))
                    for dep in deps_alese:
                        cursor.execute("INSERT INTO angajat_departament (id_angajat, departament) VALUES (%s, %s)", (nou_id, dep))
                    con.commit()
                    st.success("Adăugat!")
                    time.sleep(1)
                    st.rerun()

        cursor.close()
        con.close()

    except Exception as e:
        st.error(f"Eroare: {e}")
